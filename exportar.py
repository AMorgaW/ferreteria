# -*- coding: utf-8 -*-
"""
Exportación simple a Excel (.xlsx) compatible con Microsoft Excel.
Sin formatos visuales complejos: encabezados en negrita y ancho automático.

Los datasets monetarios e inventario salen de las mismas reglas 4A
(ReportesService / Decimal). float no es autoridad.
"""
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from services.caja_service import money


def exportar_a_excel(path: str, headers: list, rows: list, sheet_name: str = "Datos") -> str:
    """Escribe una hoja con encabezados + filas. Devuelve la ruta."""
    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_name or "Datos")[:31]

    ws.append(list(headers))
    for celda in ws[1]:
        celda.font = Font(bold=True)

    for fila in rows:
        ws.append(list(fila))

    for col_idx, _h in enumerate(headers, start=1):
        max_len = len(str(headers[col_idx - 1]))
        for fila in rows:
            if col_idx - 1 < len(fila):
                max_len = max(max_len, len(str(fila[col_idx - 1])))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 50)

    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def _money(v):
    if v is None or v == "":
        return Decimal("0.00")
    if isinstance(v, Decimal):
        return money(v)
    return money(v)


def _qty_cell(value):
    if value is None or value == "":
        return Decimal("0")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, float):
        return Decimal(str(value))
    return Decimal(str(value))


def _inventory_rows(source) -> list:
    if source is None:
        return []
    if hasattr(source, "reporte_inventario_actual"):
        return source.reporte_inventario_actual()
    if hasattr(source, "db"):
        from services.reportes_service import ReportesService

        return ReportesService(source.db).reporte_inventario_actual()
    if isinstance(source, list):
        if any(
            "cantidad_actual" not in item
            or "quantity_source" not in item
            or "valor_inventario" not in item
            for item in source
        ):
            raise ValueError(
                "El XLSX de inventario requiere el dataset canónico de ReportesService"
            )
        return source
    raise TypeError("Fuente de inventario no compatible con ReportesService")


def exportar_inventario(source, path: str) -> str:
    """Exporta inventario canónico 4A. No usa productos.stock como autoridad."""
    productos = _inventory_rows(source)
    headers = [
        "SKU",
        "Nombre",
        "Categoría",
        "Marca",
        "Presentación",
        "P. Compra",
        "P. Venta",
        "Stock",
        "Stock Mín.",
        "U. Medida",
        "Valor (compra)",
        "Activo",
    ]
    from formato import formatear_stock

    rows = []
    for item in productos:
        stock = _qty_cell(item["cantidad_actual"])
        compra = _money(item.get("precio_compra"))
        permite = item.get("permite_decimales")
        valor = _money(item["valor_inventario"])
        rows.append(
            [
                item.get("codigo_barras") or "",
                item.get("nombre") or "",
                item.get("categoria") or "",
                item.get("marca") or "",
                item.get("presentacion") or "",
                compra,
                _money(item.get("precio_venta")),
                formatear_stock(stock, permite),
                formatear_stock(item.get("stock_minimo"), permite),
                item.get("unidad_base") or item.get("unidad_medida") or "",
                valor,
                "Sí" if item.get("activo", 1) else "No",
            ]
        )
    return exportar_a_excel(path, headers, rows, "Inventario")


def exportar_ventas(ventas: list, path: str) -> str:
    """Exporta una lista de ventas (dicts de ReportesService)."""
    headers = [
        "Factura",
        "Fecha",
        "Cliente",
        "Vendedor",
        "Subtotal",
        "Descuento",
        "Total",
        "Método de Pago",
        "Estado",
    ]
    rows = []
    for venta in ventas:
        rows.append(
            [
                venta.get("numero_factura") or "",
                str(venta.get("fecha") or ""),
                venta.get("cliente_nombre") or "Consumidor Final",
                venta.get("vendedor") or "",
                _money(venta.get("subtotal")),
                _money(venta.get("descuento")),
                _money(venta.get("total")),
                venta.get("metodo_pago") or "",
                venta.get("estado") or "",
            ]
        )
    return exportar_a_excel(path, headers, rows, "Ventas")


def exportar_compras(compras: list, path: str) -> str:
    """Exporta una lista de compras (dicts)."""
    headers = [
        "ID",
        "Fecha",
        "Proveedor",
        "Factura",
        "Tipo",
        "Total",
        "Estado Pago",
        "Saldo",
    ]
    rows = []
    for compra in compras:
        rows.append(
            [
                compra.get("id") or compra.get("numero") or "",
                str(compra.get("fecha") or ""),
                compra.get("proveedor_nombre") or compra.get("proveedor") or "",
                compra.get("numero_factura") or compra.get("factura") or "",
                compra.get("tipo_pago") or compra.get("tipo") or "",
                _money(compra.get("total") or compra.get("costo_total")),
                compra.get("estado_pago") or compra.get("estado") or "",
                _money(compra.get("saldo_pendiente") or compra.get("saldo") or 0),
            ]
        )
    return exportar_a_excel(path, headers, rows, "Compras")


def exportar_tabla_qt(table, path: str, sheet_name: str = "Reporte") -> str:
    """Exporta el contenido visible de un QTableWidget a Excel (duck-typed,
    no requiere importar PySide6 aquí). Sirve para cualquier reporte tabular."""
    cols = table.columnCount()
    headers = []
    for col in range(cols):
        item = table.horizontalHeaderItem(col)
        headers.append(item.text() if item is not None else f"Col{col + 1}")
    rows = []
    for row in range(table.rowCount()):
        if table.isRowHidden(row):
            continue
        fila = []
        for col in range(cols):
            item = table.item(row, col)
            fila.append(item.text() if item is not None else "")
        rows.append(fila)
    return exportar_a_excel(path, headers, rows, sheet_name)


def nombre_sugerido(prefijo: str) -> str:
    """Nombre de archivo sugerido con marca temporal."""
    return f"{prefijo}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
