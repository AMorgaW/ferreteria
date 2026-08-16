# -*- coding: utf-8 -*-
"""XLSX parity, UI consume service, local-first."""
from __future__ import annotations

import inspect
import os
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest import mock

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from tests.fase4a.helpers import (
    complete_sale,
    ensure_inventory_balance,
    phase4a_env,
    reportes_service,
    seed_pos_product,
)


class ExportsUiReportingTest(unittest.TestCase):
    def test_20_xlsx_inventory_parity(self):
        import exportar
        from openpyxl import load_workbook

        with phase4a_env() as env:
            pid, lid, _p = seed_pos_product(env, stock=99, name="Broca")
            conn = env.connect()
            try:
                conn.execute(
                    "UPDATE productos SET precio_compra=250, stock_minimo=1 WHERE id=?",
                    (pid,),
                )
                ensure_inventory_balance(conn, lid, 7000)
                conn.commit()
            finally:
                conn.close()
            svc = reportes_service(env)
            dataset = [row for row in svc.reporte_inventario_actual() if row["id"] == pid][0]
            with tempfile.TemporaryDirectory() as tmp:
                path = str(Path(tmp) / "inv.xlsx")
                exportar.exportar_inventario(svc, path)
                wb = load_workbook(path)
                ws = wb.active
                found = None
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if row[1] == "Broca":
                        found = row
                        break
                self.assertIsNotNone(found)
                valor = Decimal(str(found[10]))
                self.assertEqual(valor, dataset["valor_inventario"])
                self.assertEqual(dataset["cantidad_actual"], Decimal("7"))
                self.assertNotEqual(dataset["cantidad_actual"], Decimal("99"))

    def test_21_xlsx_monetary_parity(self):
        import exportar
        from openpyxl import load_workbook

        with phase4a_env() as env:
            seed_pos_product(env, stock=20)
            venta = complete_sale(
                env, [{"producto_id": 1, "cantidad": 2, "precio_unitario": 15}]
            )
            hoy = str(venta.fecha)[:10]
            svc = reportes_service(env)
            data = svc.reporte_ventas_periodo(hoy, hoy)
            with tempfile.TemporaryDirectory() as tmp:
                path = str(Path(tmp) / "ventas.xlsx")
                exportar.exportar_ventas(data["ventas"], path)
                wb = load_workbook(path)
                ws = wb.active
                row = next(ws.iter_rows(min_row=2, values_only=True))
                total = Decimal(str(row[6]))
                self.assertEqual(total, data["ventas"][0]["total"])
                self.assertEqual(total, Decimal("30.00"))
                self.assertEqual(data["resumen"]["gross_sales"], Decimal("30.00"))

    def test_23_ui_consumes_service(self):
        import ui.reportes_ui as reportes_ui

        source = inspect.getsource(reportes_ui)
        self.assertNotIn("SELECT ", source)
        self.assertIn("reporte_dia", source)
        self.assertIn("detalle_venta_lineas", source)
        self.assertIn("reportes_service.", source)

    def test_22_export_local_first(self):
        import exportar

        with phase4a_env() as env:
            pid, lid, _p = seed_pos_product(env, stock=1, name="Tuerca")
            conn = env.connect()
            try:
                ensure_inventory_balance(conn, lid, 8000)
                conn.commit()
            finally:
                conn.close()

            def boom(*_a, **_k):
                raise AssertionError("network")

            with mock.patch("psycopg2.connect", side_effect=boom):
                with tempfile.TemporaryDirectory() as tmp:
                    path = str(Path(tmp) / "inv.xlsx")
                    exportar.exportar_inventario(reportes_service(env), path)
                    self.assertTrue(Path(path).exists())

    def test_25_inventory_export_rejects_legacy_stock_list(self):
        import exportar

        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "legacy.xlsx")
            with self.assertRaisesRegex(ValueError, "dataset canónico"):
                exportar.exportar_inventario(
                    [{"nombre": "Legacy", "stock": 99, "precio_compra": 10}],
                    path,
                )


if __name__ == "__main__":
    unittest.main(verbosity=2)
