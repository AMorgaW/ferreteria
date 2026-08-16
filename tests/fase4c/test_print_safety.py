# -*- coding: utf-8 -*-
"""PDF/print safety, local-first, XLSX regression, UI consume service, DB real."""
from __future__ import annotations

import inspect
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from services.document_service import (
    render_operational_text,
    suggested_pdf_filename,
)
from tests.fase4a.helpers import reportes_service
from tests.fase4c.helpers import (
    assert_commercial_untouched,
    business_fingerprint,
    commercial_db_mtime,
    complete_sale,
    docs,
    phase4c_env,
    seed_pos_product,
)


class PrintPdfSafetyTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from PySide6.QtWidgets import QApplication

        cls.app = QApplication.instance() or QApplication([])

    def test_05_18_19_print_pdf_zero_writes(self):
        from ui.imprimir_factura import print_operational_text, save_operational_pdf

        with phase4c_env() as env:
            seed_pos_product(env, stock=5, name="Tuerca")
            venta = complete_sale(
                env, [{"producto_id": 1, "cantidad": 1, "precio_unitario": 15}]
            )
            document = docs(env).load_sale(venta.id)
            before = business_fingerprint(env)
            text = render_operational_text(document)
            cancelled = print_operational_text(
                text, cancelled=True, show_dialogs=False
            )
            failed = print_operational_text(
                text, available_printers=[], show_dialogs=False
            )
            with mock.patch(
                "ui.imprimir_factura._imprimir_directo",
                side_effect=RuntimeError("printer down"),
            ):
                errored = print_operational_text(
                    text,
                    available_printers=["FakePrinter"],
                    printer_name="FakePrinter",
                    show_dialogs=False,
                )
            self.assertEqual(cancelled, "cancelled")
            self.assertEqual(failed, "no_printer")
            self.assertTrue(errored.startswith("error:"))
            name = suggested_pdf_filename(document)
            self.assertTrue(name.startswith("VENTA_"))
            self.assertTrue(name.endswith(".pdf"))
            self.assertNotRegex(name, r'[<>:"/\\|?*]')
            with tempfile.TemporaryDirectory() as tmp:
                path = str(Path(tmp) / name)
                save_operational_pdf(document, path)
                self.assertTrue(Path(path).exists())
                self.assertGreater(Path(path).stat().st_size, 0)
            self.assertEqual(before, business_fingerprint(env))

    def test_20_local_first_no_remote(self):
        with phase4c_env() as env:
            seed_pos_product(env, stock=3)
            venta = complete_sale(
                env, [{"producto_id": 1, "cantidad": 1, "precio_unitario": 8}]
            )

            def boom(*_a, **_k):
                raise AssertionError("network")

            with mock.patch("psycopg2.connect", side_effect=boom):
                doc = docs(env).load_sale(venta.id)
                text = render_operational_text(doc)
            self.assertIn("DOCUMENTO OPERACIONAL", text)

    def test_22_xlsx_regression(self):
        import exportar
        from openpyxl import load_workbook

        with phase4c_env() as env:
            seed_pos_product(env, stock=20)
            venta = complete_sale(
                env, [{"producto_id": 1, "cantidad": 2, "precio_unitario": 15}]
            )
            hoy = str(venta.fecha)[:10]
            data = reportes_service(env).reporte_ventas_periodo(hoy, hoy)
            with tempfile.TemporaryDirectory() as tmp:
                path = str(Path(tmp) / "ventas.xlsx")
                exportar.exportar_ventas(data["ventas"], path)
                wb = load_workbook(path)
                row = next(wb.active.iter_rows(min_row=2, values_only=True))
                from decimal import Decimal

                self.assertEqual(Decimal(str(row[6])), Decimal("30.00"))

    def test_23_ui_uses_document_service(self):
        import ui.caja_ui as caja_ui
        import ui.compras_ui as compras_ui
        import ui.imprimir_factura as imprimir
        import ui.movimientos_ui as movimientos_ui
        import ui.returns_ui as returns_ui
        import ui.ventas_ui_modern as ventas_ui

        self.assertIn("DocumentService", inspect.getsource(imprimir))
        self.assertIn("imprimir_documento", inspect.getsource(imprimir))
        self.assertIn("db=self.db_manager", inspect.getsource(ventas_ui))
        self.assertIn("imprimir_por_identidad", inspect.getsource(movimientos_ui))
        self.assertIn("imprimir_por_identidad", inspect.getsource(compras_ui))
        self.assertIn("imprimir_por_identidad", inspect.getsource(returns_ui))
        self.assertIn("CASH_CLOSE", inspect.getsource(caja_ui))
        body = inspect.getsource(imprimir._generar_contenido_factura)
        self.assertIn("sale_document_from_payload", body)
        self.assertIn("render_operational_text", body)

    def test_24_commercial_db_untouched(self):
        before = commercial_db_mtime()
        with phase4c_env() as env:
            seed_pos_product(env, stock=2)
            venta = complete_sale(
                env, [{"producto_id": 1, "cantidad": 1, "precio_unitario": 5}]
            )
            docs(env).load_sale(venta.id)
        assert_commercial_untouched(before)


if __name__ == "__main__":
    unittest.main(verbosity=2)
