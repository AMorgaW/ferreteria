from __future__ import annotations

import hashlib
import sqlite3
import tempfile
import unittest
import uuid
from pathlib import Path

from openpyxl import load_workbook

from product_inventory_contract import (
    DATA_START_ROW,
    FIELD_BY_KEY,
    PRODUCT_FIELDS,
    PRODUCT_HEADERS,
    REQUIRED_READY_KEYS,
)
from scripts.generate_inventory_workbook import generate_workbook
from scripts.validate_inventory_workbook import validate_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def col(key: str) -> int:
    return next(index for index, field in enumerate(PRODUCT_FIELDS, 1) if field.key == key)


def create_legacy_db(path: Path) -> None:
    conn = sqlite3.connect(path)
    try:
        conn.executescript(
            """
            CREATE TABLE proveedores (
                id INTEGER PRIMARY KEY, nombre TEXT, activo INTEGER DEFAULT 1
            );
            CREATE TABLE productos (
                id INTEGER PRIMARY KEY,
                local_id TEXT UNIQUE,
                codigo_barras TEXT UNIQUE,
                nombre TEXT NOT NULL,
                categoria TEXT,
                marca TEXT,
                presentacion TEXT,
                proveedor_id INTEGER,
                precio_compra REAL DEFAULT 0,
                precio_venta REAL NOT NULL,
                stock_minimo INTEGER DEFAULT 10,
                unidad_medida TEXT DEFAULT 'UNIDAD',
                permite_decimales INTEGER DEFAULT 0,
                viene_en_caja INTEGER DEFAULT 0,
                unidades_por_caja INTEGER DEFAULT 1,
                vende_por_empaque INTEGER DEFAULT 0,
                ubicacion TEXT,
                descripcion TEXT,
                stock REAL DEFAULT 0,
                activo INTEGER DEFAULT 1
            );
            INSERT INTO proveedores(id, nombre, activo) VALUES (1, 'Proveedor Uno', 1);
            """
        )
        conn.execute(
            "INSERT INTO productos VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (
                1, "11111111-1111-4111-8111-111111111111", "SKU-LEGACY-1",
                "Pintura blanca", "Pinturas", "Marca Uno", "1 galón", 1,
                50000, 70000, 10, "GALÓN", 1, 0, 1, 0, "A-1", "Mate", 5.125, 1,
            ),
        )
        conn.execute(
            "INSERT INTO productos VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (
                2, "22222222-2222-4222-8222-222222222222", None,
                "Martillo", "Herramientas Manuales", "Marca Dos", None, None,
                20000, 30000, 5, "UNIDAD", 0, 0, 1, 0, None, None, 3, 1,
            ),
        )
        conn.commit()
    finally:
        conn.close()


def make_ready(ws, row: int, *, name: str = "Producto listo") -> None:
    values = {
        "nombre": name,
        "categoria": "Pinturas",
        "precio_compra": 100,
        "precio_venta": 150,
        "stock_minimo": 10,
        "unidad_medida": "UNIDAD",
        "permite_decimales": "NO",
        "viene_en_caja": "NO",
        "unidades_por_caja": 1,
        "vende_por_empaque": "NO",
        "cantidad_contada": 2,
        "estado_barcode": "PENDIENTE",
    }
    for key, value in values.items():
        ws.cell(row, col(key), value)


class ProductContractTest(unittest.TestCase):
    def test_01_contract_discovery_matches_repository_and_ui(self):
        repo = (REPO_ROOT / "repositories" / "productos_repo.py").read_text(encoding="utf-8")
        ui = (REPO_ROOT / "ui" / "productos_ui.py").read_text(encoding="utf-8")
        self.assertIn("producto.validar()", repo)
        self.assertIn("mismo nombre + marca + presentación", repo)
        for label in ("Nombre del Producto *", "Categoría *", "Precio de Compra *", "Precio de Venta *", "Unidad de Medida Base *", "Stock Mínimo *"):
            self.assertIn(label, ui)

    def test_02_required_fields_are_canonical(self):
        self.assertEqual(
            REQUIRED_READY_KEYS,
            ("nombre", "categoria", "precio_compra", "precio_venta", "stock_minimo", "unidad_medida", "cantidad_contada"),
        )
        self.assertNotIn("codigo_barras", REQUIRED_READY_KEYS)

    def test_03_legacy_only_fields_not_forced_into_products(self):
        self.assertFalse(FIELD_BY_KEY["precio_mayorista"].include_in_products)
        self.assertFalse(FIELD_BY_KEY["stock_maximo"].include_in_products)
        self.assertEqual(FIELD_BY_KEY["codigo_barras"].classification, "BARCODE_PENDING")


class WorkbookGenerationTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def blank(self, rows: int = 3) -> Path:
        path = self.root / "blank.xlsx"
        generate_workbook(path, blank_rows=rows)
        return path

    def test_04_blank_generation_has_expected_sheets(self):
        path = self.blank()
        wb = load_workbook(path, data_only=False)
        try:
            self.assertEqual(wb.sheetnames, ["INSTRUCCIONES", "PRODUCTOS", "CATALOGOS", "RESUMEN", "MAPEO_FERREPRO", "BARCODES_PENDIENTES"])
        finally:
            wb.close()

    def test_05_required_fields_are_reflected_in_workbook(self):
        wb = load_workbook(self.blank(), data_only=False)
        try:
            headers = tuple(wb["PRODUCTOS"].cell(4, i).value for i in range(1, len(PRODUCT_HEADERS) + 1))
            self.assertEqual(headers, PRODUCT_HEADERS)
            for key in REQUIRED_READY_KEYS:
                self.assertEqual(FIELD_BY_KEY[key].classification, "INVENTORY_COUNT_DATA" if key == "cantidad_contada" else "REQUIRED_FOR_NEW_PRODUCT")
        finally:
            wb.close()

    def test_06_barcode_not_required_for_ready_without_barcode(self):
        path = self.blank()
        wb = load_workbook(path)
        make_ready(wb["PRODUCTOS"], 5)
        wb.save(path)
        wb.close()
        report = validate_workbook(path)
        self.assertEqual(report["row_status"]["5"], "LISTO_SIN_BARCODE")

    def test_07_staging_id_is_stable_after_save_and_reopen(self):
        path = self.blank()
        wb = load_workbook(path)
        value = wb["PRODUCTOS"].cell(5, col("registro_inventario_id")).value
        wb.save(path)
        wb.close()
        wb = load_workbook(path)
        try:
            self.assertEqual(wb["PRODUCTOS"].cell(5, col("registro_inventario_id")).value, value)
            uuid.UUID(value)
        finally:
            wb.close()

    def test_08_formulas_exist_for_inventory_calculations(self):
        wb = load_workbook(self.blank(), data_only=False)
        try:
            ws = wb["PRODUCTOS"]
            for key in ("diferencia_unidades", "valor_inventario_costo", "valor_potencial_venta", "margen_unitario", "margen_porcentaje", "estado_registro", "errores_validacion"):
                self.assertTrue(str(ws.cell(5, col(key)).value).startswith("="), key)
        finally:
            wb.close()

    def test_09_mapping_sheet_covers_every_product_column(self):
        wb = load_workbook(self.blank(), data_only=False)
        try:
            mapped = {wb["MAPEO_FERREPRO"].cell(row, 1).value for row in range(5, wb["MAPEO_FERREPRO"].max_row + 1)}
            self.assertTrue(set(PRODUCT_HEADERS).issubset(mapped))
        finally:
            wb.close()

    def test_10_data_validations_are_present(self):
        wb = load_workbook(self.blank(), data_only=False)
        try:
            self.assertGreaterEqual(len(wb["PRODUCTOS"].data_validations.dataValidation), 8)
        finally:
            wb.close()

    def test_10b_summary_does_not_count_reserved_new_slots(self):
        wb = load_workbook(self.blank(), data_only=False)
        try:
            formulas = [wb["RESUMEN"].cell(row, 2).value for row in range(4, 15)]
            self.assertTrue(any("$G$" in str(value) and '"NUEVO"' in str(value) for value in formulas))
            self.assertFalse(any("$AH$" in str(value) and '"NUEVO"' in str(value) for value in formulas))
        finally:
            wb.close()

    def test_10c_barcode_view_formulas_are_error_guarded(self):
        wb = load_workbook(self.blank(), data_only=False)
        try:
            for cell in wb["BARCODES_PENDIENTES"][5]:
                self.assertIn("IFERROR", str(cell.value))
        finally:
            wb.close()

    def test_11_from_db_preloads_existing_and_leaves_count_blank(self):
        db = self.root / "legacy.db"
        create_legacy_db(db)
        out = self.root / "from-db.xlsx"
        result = generate_workbook(out, db_path=db, blank_rows=2)
        self.assertEqual(result["existing_products"], 2)
        wb = load_workbook(out)
        try:
            ws = wb["PRODUCTOS"]
            self.assertEqual(ws.cell(5, col("nombre")).value, "Pintura blanca")
            self.assertIsNone(ws.cell(5, col("cantidad_contada")).value)
            self.assertEqual(ws.cell(5, col("codigo_sistema_actual")).value, "SKU-LEGACY-1")
            self.assertIsNone(ws.cell(5, col("codigo_barras")).value)
            self.assertEqual(ws.cell(5, col("estado_barcode")).value, "PENDIENTE")
        finally:
            wb.close()

    def test_12_from_db_is_read_only(self):
        db = self.root / "legacy.db"
        create_legacy_db(db)
        before = (sha256(db), db.stat().st_size, db.stat().st_mtime_ns)
        generate_workbook(self.root / "from-db.xlsx", db_path=db, blank_rows=2)
        after = (sha256(db), db.stat().st_size, db.stat().st_mtime_ns)
        self.assertEqual(before, after)

    def test_13_existing_staging_id_reproducible_between_generations(self):
        db = self.root / "legacy.db"
        create_legacy_db(db)
        first, second = self.root / "a.xlsx", self.root / "b.xlsx"
        generate_workbook(first, db_path=db, blank_rows=1)
        generate_workbook(second, db_path=db, blank_rows=1)
        wa, wb = load_workbook(first), load_workbook(second)
        try:
            self.assertEqual(wa["PRODUCTOS"].cell(5, 1).value, wb["PRODUCTOS"].cell(5, 1).value)
        finally:
            wa.close(); wb.close()


class WorkbookValidationTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "inventory.xlsx"
        generate_workbook(self.path, blank_rows=3)

    def tearDown(self):
        self.tmp.cleanup()

    def edit(self, callback) -> None:
        wb = load_workbook(self.path)
        callback(wb["PRODUCTOS"])
        wb.save(self.path)
        wb.close()

    def test_14_invalid_price_detected(self):
        self.edit(lambda ws: (make_ready(ws, 5), ws.cell(5, col("precio_venta"), -1)))
        codes = {item["code"] for item in validate_workbook(self.path)["issues"]}
        self.assertIn("INVALID_PRICE", codes)

    def test_15_invalid_quantity_detected(self):
        self.edit(lambda ws: (make_ready(ws, 5), ws.cell(5, col("cantidad_contada"), "abc")))
        codes = {item["code"] for item in validate_workbook(self.path)["issues"]}
        self.assertIn("INVALID_NUMBER", codes)

    def test_16_more_than_three_quantity_decimals_rejected(self):
        self.edit(lambda ws: (make_ready(ws, 5), ws.cell(5, col("cantidad_contada"), 1.2345), ws.cell(5, col("permite_decimales"), "SI")))
        codes = {item["code"] for item in validate_workbook(self.path)["issues"]}
        self.assertIn("QUANTITY_SCALE", codes)

    def test_17_duplicate_staging_id_rejected(self):
        def mutate(ws):
            ws.cell(6, col("registro_inventario_id"), ws.cell(5, col("registro_inventario_id")).value)
        self.edit(mutate)
        codes = {item["code"] for item in validate_workbook(self.path)["issues"]}
        self.assertIn("STAGING_ID_DUPLICATE", codes)

    def test_18_potential_duplicate_product_flagged(self):
        self.edit(lambda ws: (make_ready(ws, 5, name="Mismo"), make_ready(ws, 6, name="Mismo")))
        report = validate_workbook(self.path)
        self.assertIn("POTENTIAL_DUPLICATE", {item["code"] for item in report["issues"]})
        self.assertEqual(report["row_status"]["5"], "REVISAR")

    def test_19_duplicate_barcode_flagged(self):
        def mutate(ws):
            for row in (5, 6):
                make_ready(ws, row, name=f"Producto {row}")
                ws.cell(row, col("codigo_barras"), "7701234567890")
                ws.cell(row, col("estado_barcode"), "ESCANEADO")
        self.edit(mutate)
        self.assertIn("BARCODE_DUPLICATE", {item["code"] for item in validate_workbook(self.path)["issues"]})

    def test_20_validator_does_not_mutate_workbook(self):
        before = (sha256(self.path), self.path.stat().st_size, self.path.stat().st_mtime_ns)
        validate_workbook(self.path)
        after = (sha256(self.path), self.path.stat().st_size, self.path.stat().st_mtime_ns)
        self.assertEqual(before, after)

    def test_21_status_ready_complete_after_barcode(self):
        def mutate(ws):
            make_ready(ws, 5)
            ws.cell(5, col("codigo_barras"), "7701234567890")
            ws.cell(5, col("estado_barcode"), "ESCANEADO")
        self.edit(mutate)
        self.assertEqual(validate_workbook(self.path)["row_status"]["5"], "LISTO_COMPLETO")

    def test_22_header_tampering_is_structural_error(self):
        self.edit(lambda ws: ws.cell(4, 1, "ID_CAMBIADO"))
        self.assertTrue(validate_workbook(self.path)["structural_errors"])


if __name__ == "__main__":
    unittest.main()
