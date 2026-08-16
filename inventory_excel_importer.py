# -*- coding: utf-8 -*-
"""Parser seguro y testeable del XLSX operacional de inventario físico.

No conoce Qt, no abre conexiones y no muta catálogo, stock ni barcodes. Toda
fila se normaliza y valida antes de que el servicio decida persistir staging.
"""
from __future__ import annotations

import hashlib
import json
import re
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Mapping, Optional, Sequence

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from product_inventory_contract import (
    CATEGORY_UI_VALUES,
    NO_PACKAGE_VALUE,
    OPERATIONAL_HEADER_TO_KEY,
    OPERATIONAL_INVENTORY_HEADERS,
    OPERATIONAL_REQUIRED_KEYS,
    PACKAGE_PRESENTATION_VALUES,
    UNIT_UI_VALUES,
    has_max_three_decimals,
    normalized_text,
)


MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_COMPRESSION_RATIO = 100
MAX_ROWS = 10_000
MAX_COLUMNS = 100
MAX_STRING_LENGTH = 500
HEADER_SEARCH_ROWS = 20
QUANTITY_SCALE = Decimal("1000")


class InventoryWorkbookError(ValueError):
    """El archivo completo es inseguro o no cumple el contrato."""


@dataclass(frozen=True)
class ImportIssue:
    severity: str
    code: str
    field: str
    message: str

    def as_dict(self) -> dict:
        return {
            "severity": self.severity,
            "code": self.code,
            "field": self.field,
            "message": self.message,
        }


@dataclass
class ParsedInventoryRow:
    excel_row_number: int
    values: dict[str, Any]
    raw_payload: dict[str, Any]
    quantity_counted: Optional[Decimal]
    quantity_counted_scaled: Optional[int]
    barcode_candidate: Optional[str]
    barcode_status: str
    validation_status: str
    errors: list[ImportIssue] = field(default_factory=list)
    warnings: list[ImportIssue] = field(default_factory=list)
    infos: list[ImportIssue] = field(default_factory=list)
    row_hash: str = ""
    duplicate_candidate: bool = False

    @property
    def all_issues(self) -> list[ImportIssue]:
        return [*self.errors, *self.warnings, *self.infos]


@dataclass(frozen=True)
class ParsedInventoryWorkbook:
    source_path: Path
    source_filename: str
    source_sha256: str
    sheet_name: str
    header_row: int
    rows: tuple[ParsedInventoryRow, ...]
    extra_headers: tuple[str, ...]
    workbook_warnings: tuple[ImportIssue, ...]


def _header_token(value: Any) -> str:
    text = normalized_text(value)
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


HEADER_TOKEN_TO_KEY = {
    _header_token(header): key for header, key in OPERATIONAL_HEADER_TO_KEY.items()
}


def _canonical_decimal(value: Decimal) -> str:
    text = format(value, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def _canonical_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return _canonical_decimal(value)
    if value is None or isinstance(value, (str, int, bool)):
        return value
    return str(value)


def canonical_row_hash(values: Mapping[str, Any]) -> str:
    material = {
        key: _canonical_value(values.get(key))
        for key in sorted(OPERATIONAL_HEADER_TO_KEY.values())
    }
    encoded = json.dumps(
        material, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def decimal_to_scaled(value: Decimal) -> int:
    if not has_max_three_decimals(value):
        raise InvalidOperation("cantidad excede 3 decimales")
    return int((value * QUANTITY_SCALE).to_integral_exact())


def parse_decimal_input(value: Any, *, money: bool = False) -> Optional[Decimal]:
    """Convierte input XLSX a Decimal sin usar float como representación final."""
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise InvalidOperation("boolean no es número")
    if isinstance(value, Decimal):
        result = value
    elif isinstance(value, (int, float)):
        result = Decimal(str(value))
    elif isinstance(value, str):
        text = value.strip()
        if money:
            text = re.sub(r"(?i)\bCOP\b", "", text)
            text = text.replace("$", "").replace("\u00a0", "").replace(" ", "")
            if "," in text and "." in text:
                if text.rfind(",") > text.rfind("."):
                    text = text.replace(".", "").replace(",", ".")
                else:
                    text = text.replace(",", "")
            elif "," in text:
                right = text.rsplit(",", 1)[1]
                text = text.replace(",", ".") if len(right) <= 2 else text.replace(",", "")
            elif "." in text and text.count(".") == 1 and len(text.rsplit(".", 1)[1]) == 3:
                # Notación monetaria habitual COP: $ 10.000.
                text = text.replace(".", "")
            elif text.count(".") > 1:
                text = text.replace(".", "")
        try:
            result = Decimal(text)
        except InvalidOperation as exc:
            raise InvalidOperation("valor no numérico") from exc
    else:
        raise InvalidOperation("tipo numérico no admitido")
    if not result.is_finite():
        raise InvalidOperation("número no finito")
    return result


def _display_value(value: Any) -> Any:
    if isinstance(value, (str, int, bool)) or value is None:
        return value
    if isinstance(value, float):
        return str(Decimal(str(value)))
    return str(value)


def _barcode_cell_value(cell) -> str:
    value = cell.value
    if value in (None, ""):
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        fmt = str(cell.number_format or "")
        if fmt and set(fmt) == {"0"}:
            return f"{value:0{len(fmt)}d}"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _issue(target: list[ImportIssue], severity: str, code: str, field: str, message: str) -> None:
    target.append(ImportIssue(severity, code, field, message))


def _validate_zip(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if not infos or "[Content_Types].xml" not in archive.namelist():
                raise InventoryWorkbookError("XLSX inválido: estructura OOXML ausente")
            total = sum(item.file_size for item in infos)
            compressed = sum(max(item.compress_size, 1) for item in infos)
            if total > MAX_UNCOMPRESSED_BYTES:
                raise InventoryWorkbookError("XLSX excede el tamaño descomprimido permitido")
            if total / compressed > MAX_COMPRESSION_RATIO:
                raise InventoryWorkbookError("XLSX rechazado por relación de compresión insegura")
            if any(name.lower().endswith("vbaproject.bin") for name in archive.namelist()):
                raise InventoryWorkbookError("XLSX contiene macros no permitidas")
    except zipfile.BadZipFile as exc:
        raise InventoryWorkbookError("XLSX malformado") from exc


class InventoryExcelImporter:
    def __init__(
        self,
        *,
        categories: Iterable[str] = CATEGORY_UI_VALUES,
        units: Iterable[str] = UNIT_UI_VALUES,
        presentations: Iterable[str] = PACKAGE_PRESENTATION_VALUES,
        max_rows: int = MAX_ROWS,
        max_file_bytes: int = MAX_FILE_BYTES,
    ) -> None:
        self.categories = {normalized_text(item) for item in categories}
        self.units = {normalized_text(item) for item in units}
        self.presentations = {normalized_text(item) for item in presentations}
        self.max_rows = int(max_rows)
        self.max_file_bytes = int(max_file_bytes)

    def parse(self, source: str | Path) -> ParsedInventoryWorkbook:
        path = Path(source)
        if path.suffix.lower() != ".xlsx":
            raise InventoryWorkbookError("Solo se admite extensión .xlsx")
        if not path.is_file():
            raise InventoryWorkbookError("Archivo XLSX no encontrado")
        if path.stat().st_size > self.max_file_bytes:
            raise InventoryWorkbookError("XLSX excede el tamaño máximo permitido")
        _validate_zip(path)
        sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
        try:
            workbook = load_workbook(
                path, read_only=True, data_only=True, keep_links=False
            )
        except (InvalidFileException, OSError, ValueError, KeyError) as exc:
            raise InventoryWorkbookError(f"No fue posible leer el XLSX: {exc}") from exc
        try:
            sheet, header_row, columns, extras = self._locate_contract(workbook)
            rows = self._read_rows(sheet, header_row, columns)
        finally:
            workbook.close()
        warnings = tuple(
            ImportIssue("WARNING", "EXTRA_COLUMN", "header", f"Columna ignorada: {name}")
            for name in extras
        )
        return ParsedInventoryWorkbook(
            source_path=path,
            source_filename=path.name,
            source_sha256=sha256,
            sheet_name=sheet.title,
            header_row=header_row,
            rows=tuple(rows),
            extra_headers=tuple(extras),
            workbook_warnings=warnings,
        )

    def _locate_contract(self, workbook):
        best = None
        required_tokens = set(HEADER_TOKEN_TO_KEY)
        for sheet in workbook.worksheets:
            for row_number in range(1, min(sheet.max_row, HEADER_SEARCH_ROWS) + 1):
                tokens = [_header_token(sheet.cell(row_number, col).value) for col in range(1, min(sheet.max_column, MAX_COLUMNS) + 1)]
                hits = len(required_tokens.intersection(token for token in tokens if token))
                if best is None or hits > best[0]:
                    best = (hits, sheet, row_number, tokens)
        if best is None:
            raise InventoryWorkbookError("Workbook sin hojas legibles")
        _, sheet, row_number, tokens = best
        columns: dict[str, int] = {}
        extras: list[str] = []
        for index, token in enumerate(tokens, start=1):
            raw = sheet.cell(row_number, index).value
            if not token:
                continue
            key = HEADER_TOKEN_TO_KEY.get(token)
            if key:
                if key in columns:
                    raise InventoryWorkbookError(f"Encabezado duplicado: {raw}")
                columns[key] = index
            else:
                extras.append(str(raw))
        missing = [
            label for label in OPERATIONAL_INVENTORY_HEADERS
            if OPERATIONAL_HEADER_TO_KEY[label] not in columns
        ]
        if missing:
            raise InventoryWorkbookError(
                "Faltan encabezados requeridos: " + "; ".join(missing)
            )
        return sheet, row_number, columns, extras

    def _read_rows(self, sheet, header_row: int, columns: Mapping[str, int]) -> list[ParsedInventoryRow]:
        rows: list[ParsedInventoryRow] = []
        last_possible = sheet.max_row
        if last_possible - header_row > self.max_rows:
            # Plantillas pueden traer formato muy por debajo de sus datos. Solo
            # se rechaza si hay contenido real más allá del límite.
            probe_start = header_row + self.max_rows + 1
            for row_num in range(probe_start, last_possible + 1):
                if any(sheet.cell(row_num, col).value not in (None, "") for col in columns.values()):
                    raise InventoryWorkbookError(f"XLSX excede {self.max_rows} filas de datos")
        stop = min(last_possible, header_row + self.max_rows)
        for row_num in range(header_row + 1, stop + 1):
            cells = {key: sheet.cell(row_num, col) for key, col in columns.items()}
            if all(cell.value in (None, "") for cell in cells.values()):
                continue
            rows.append(self._parse_row(row_num, cells))
        self._mark_duplicates(rows)
        return rows

    def _parse_row(self, row_number: int, cells: Mapping[str, Any]) -> ParsedInventoryRow:
        errors: list[ImportIssue] = []
        warnings: list[ImportIssue] = []
        infos: list[ImportIssue] = []
        raw = {key: _display_value(cell.value) for key, cell in cells.items()}
        values: dict[str, Any] = {}

        for key, cell in cells.items():
            if isinstance(cell.value, str) and len(cell.value) > MAX_STRING_LENGTH:
                _issue(errors, "ERROR", "STRING_TOO_LONG", key, f"{key} excede {MAX_STRING_LENGTH} caracteres")

        for key in ("nombre", "categoria", "marca", "unidad_base", "presentacion_empaque"):
            value = str(cells[key].value or "").strip()
            if len(value) > MAX_STRING_LENGTH:
                _issue(errors, "ERROR", "STRING_TOO_LONG", key, f"{key} excede {MAX_STRING_LENGTH} caracteres")
                value = value[:MAX_STRING_LENGTH]
            values[key] = value

        for key in OPERATIONAL_REQUIRED_KEYS:
            if cells[key].value in (None, ""):
                _issue(errors, "ERROR", "REQUIRED_MISSING", key, f"Falta {key}")
        if not values["nombre"]:
            _issue(errors, "ERROR", "REQUIRED_MISSING", "nombre", "Producto/nombre es obligatorio")
        if len(values["nombre"]) > 150:
            _issue(errors, "ERROR", "NAME_TOO_LONG", "nombre", "Producto/nombre supera 150 caracteres")

        category = normalized_text(values["categoria"])
        if category and category not in self.categories:
            _issue(warnings, "WARNING", "NEW_CATEGORY_REVIEW", "categoria", "Categoría no catalogada; requiere revisión")
        unit = normalized_text(values["unidad_base"])
        if unit and unit not in self.units:
            _issue(warnings, "WARNING", "NEW_UNIT_REVIEW", "unidad_base", "Unidad no catalogada; requiere revisión")
        presentation = normalized_text(values["presentacion_empaque"])
        if presentation and presentation not in self.presentations:
            _issue(warnings, "WARNING", "NEW_PRESENTATION_REVIEW", "presentacion_empaque", "Presentación no catalogada; requiere revisión")

        for key in ("vende_unidad_base", "vende_medio_empaque", "vende_empaque_completo", "permite_decimales"):
            value = str(cells[key].value or "").strip().upper()
            values[key] = value
            if value not in ("SI", "NO"):
                _issue(errors, "ERROR", "INVALID_BOOLEAN", key, f"{key} debe ser SI o NO")

        numeric_fields = (
            "precio_compra", "precio_venta", "stock_minimo",
            "cantidad_base_por_empaque", "empaques_completos_contados",
            "unidades_sueltas_contadas", "cantidad_total_excel",
        )
        for key in numeric_fields:
            try:
                values[key] = parse_decimal_input(
                    cells[key].value, money=key in ("precio_compra", "precio_venta")
                )
            except InvalidOperation:
                values[key] = None
                _issue(errors, "ERROR", "INVALID_NUMBER", key, f"{key} debe ser numérico")

        cost, sale = values["precio_compra"], values["precio_venta"]
        if cost is not None and cost < 0:
            _issue(errors, "ERROR", "INVALID_PURCHASE_PRICE", "precio_compra", "Precio de compra debe ser >= 0")
        if sale is not None and sale <= 0:
            _issue(errors, "ERROR", "INVALID_SALE_PRICE", "precio_venta", "Precio de venta debe ser > 0")
        if cost is not None and sale is not None and sale < cost:
            _issue(errors, "ERROR", "SALE_BELOW_COST", "precio_venta", "Precio de venta no puede ser menor al costo")

        stock_min = values["stock_minimo"]
        if stock_min is not None and (stock_min < 0 or not has_max_three_decimals(stock_min)):
            _issue(errors, "ERROR", "INVALID_STOCK_MIN", "stock_minimo", "Stock mínimo debe ser >= 0 y tener máximo 3 decimales")

        package_size = values["cantidad_base_por_empaque"]
        packages = values["empaques_completos_contados"]
        loose = values["unidades_sueltas_contadas"]
        excel_total = values["cantidad_total_excel"]
        no_package = normalized_text(values["presentacion_empaque"]) == normalized_text(NO_PACKAGE_VALUE)
        for key, value in (("cantidad_base_por_empaque", package_size), ("empaques_completos_contados", packages), ("unidades_sueltas_contadas", loose), ("cantidad_total_excel", excel_total)):
            if value is not None and (value < 0 or not has_max_three_decimals(value)):
                _issue(errors, "ERROR", "INVALID_QUANTITY", key, f"{key} debe ser >= 0 y tener máximo 3 decimales")
        if no_package:
            if package_size is None:
                package_size = values["cantidad_base_por_empaque"] = Decimal("1")
            elif package_size != 1:
                _issue(errors, "ERROR", "NO_PACKAGE_SIZE", "cantidad_base_por_empaque", "SIN EMPAQUE admite vacío o 1")
            if packages not in (None, Decimal("0")):
                _issue(errors, "ERROR", "NO_PACKAGE_COMPLETE_COUNT", "empaques_completos_contados", "SIN EMPAQUE se cuenta solo en unidades base sueltas")
            expected = loose
        else:
            if package_size is None or package_size <= 0:
                _issue(errors, "ERROR", "PACKAGE_SIZE_REQUIRED", "cantidad_base_por_empaque", "La presentación requiere cantidad base por empaque > 0")
                expected = None
            else:
                expected = package_size * packages + loose if packages is not None and loose is not None else None

        if packages is not None and packages != packages.to_integral_value():
            _issue(errors, "ERROR", "COMPLETE_PACKAGES_NOT_INTEGER", "empaques_completos_contados", "Empaques completos contados debe ser entero")
        if all(values.get(key) == "NO" for key in ("vende_unidad_base", "vende_medio_empaque", "vende_empaque_completo")):
            _issue(errors, "ERROR", "NO_SALE_MODE", "vende_unidad_base", "Debe habilitarse al menos un modo de venta")
        if values["vende_medio_empaque"] == "SI":
            if no_package:
                _issue(errors, "ERROR", "HALF_PACKAGE_WITHOUT_PACKAGE", "vende_medio_empaque", "Medio empaque requiere presentación real")
            elif package_size is not None:
                half = package_size / Decimal("2")
                values["cantidad_medio_empaque"] = half
                if half != half.to_integral_value() and values["permite_decimales"] == "NO":
                    _issue(warnings, "WARNING", "REVIEW_HALF_PACKAGE_FRACTION", "vende_medio_empaque", "Medio empaque produce cantidad fraccionaria; no se redondeó")
        else:
            values["cantidad_medio_empaque"] = None

        if expected is not None and excel_total is not None and expected != excel_total:
            _issue(errors, "ERROR", "ERROR_COUNT_MISMATCH", "cantidad_total_excel", f"Total Excel {_canonical_decimal(excel_total)} difiere del recalculado {_canonical_decimal(expected)}")
        values["cantidad_total_recalculada"] = expected
        if expected is not None and values["permite_decimales"] == "NO" and expected != expected.to_integral_value():
            _issue(errors, "ERROR", "FRACTION_NOT_ALLOWED", "cantidad_total_excel", "El conteo físico es fraccionario y el producto no permite decimales")

        barcode_1 = _barcode_cell_value(cells["barcode_primero"])
        barcode_2 = _barcode_cell_value(cells["barcode_segundo"])
        if len(barcode_1) > MAX_STRING_LENGTH or len(barcode_2) > MAX_STRING_LENGTH:
            barcode_1 = barcode_1[:MAX_STRING_LENGTH]
            barcode_2 = barcode_2[:MAX_STRING_LENGTH]
        values["barcode_primero"] = barcode_1
        values["barcode_segundo"] = barcode_2
        candidate: Optional[str] = None
        if not barcode_1 and not barcode_2:
            barcode_status = "BARCODE_PENDING"
        elif barcode_1 and barcode_2 and barcode_1 == barcode_2:
            candidate = barcode_1
            barcode_status = "BARCODE_EXCEL_CANDIDATE"
            _issue(infos, "INFO", "BARCODE_REQUIRES_VERIFICATION", "barcode_primero", "Candidate Excel; no es barcode VERIFIED")
        elif barcode_1 and barcode_2:
            barcode_status = "BARCODE_ERROR"
            _issue(errors, "ERROR", "BARCODE_SCAN_MISMATCH", "barcode_segundo", "Los dos campos de barcode difieren")
        else:
            barcode_status = "BARCODE_ERROR"
            _issue(errors, "ERROR", "BARCODE_PAIR_INCOMPLETE", "barcode_segundo", "Barcode incompleto: se requieren ambos campos o ninguno")

        status = "ERROR" if errors else "WARNING" if warnings else "VALID"
        scaled = None
        if expected is not None and has_max_three_decimals(expected) and expected >= 0:
            scaled = decimal_to_scaled(expected)
        row = ParsedInventoryRow(
            excel_row_number=row_number,
            values=values,
            raw_payload=raw,
            quantity_counted=expected,
            quantity_counted_scaled=scaled,
            barcode_candidate=candidate,
            barcode_status=barcode_status,
            validation_status=status,
            errors=errors,
            warnings=warnings,
            infos=infos,
        )
        row.row_hash = canonical_row_hash(values)
        return row

    @staticmethod
    def _mark_duplicates(rows: Sequence[ParsedInventoryRow]) -> None:
        by_hash: dict[str, list[ParsedInventoryRow]] = {}
        by_barcode: dict[str, list[ParsedInventoryRow]] = {}
        by_identity: dict[tuple[str, str, str, str], list[ParsedInventoryRow]] = {}
        for row in rows:
            by_hash.setdefault(row.row_hash, []).append(row)
            if row.barcode_candidate:
                by_barcode.setdefault(row.barcode_candidate, []).append(row)
            identity = tuple(normalized_text(row.values.get(k)) for k in ("nombre", "marca", "presentacion_empaque", "categoria"))
            by_identity.setdefault(identity, []).append(row)
        duplicate_groups = [group for group in by_hash.values() if len(group) > 1]
        duplicate_groups += [group for group in by_barcode.values() if len(group) > 1]
        duplicate_groups += [group for identity, group in by_identity.items() if identity[0] and len(group) > 1]
        seen: set[tuple[int, str]] = set()
        for group in duplicate_groups:
            for row in group:
                marker = (row.excel_row_number, "DUPLICATE_CANDIDATE")
                if marker in seen:
                    continue
                seen.add(marker)
                row.duplicate_candidate = True
                row.warnings.append(ImportIssue("WARNING", "DUPLICATE_CANDIDATE", "row", "Posible producto/barcode repetido dentro del archivo; no se sumó ni eliminó"))
                if not row.errors:
                    row.validation_status = "WARNING"
