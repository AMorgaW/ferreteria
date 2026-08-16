# -*- coding: utf-8 -*-
"""Valida el XLSX de inventario sin importar ni modificar FERREPRO/SQLite."""
from __future__ import annotations

import argparse
import csv
import json
import sys
import uuid
from dataclasses import asdict
from pathlib import Path
from typing import Any, Mapping

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from product_inventory_contract import (
    DATA_START_ROW,
    HEADER_ROW,
    PRODUCT_HEADERS,
    ValidationIssue,
    is_active_record,
    normalized_text,
    status_from_issues,
    validate_record,
)

REQUIRED_SHEETS = (
    "INSTRUCCIONES", "PRODUCTOS", "CATALOGOS", "RESUMEN",
    "MAPEO_FERREPRO", "BARCODES_PENDIENTES",
)


def _catalog_values(ws, column: int) -> list[str]:
    values = []
    for row in range(5, ws.max_row + 1):
        value = ws.cell(row, column).value
        if value not in (None, ""):
            values.append(str(value).strip())
    return values


def _product_last_row(ws) -> int:
    if "ProductosInventario" in ws.tables:
        _min_col, _min_row, _max_col, max_row = range_boundaries(
            ws.tables["ProductosInventario"].ref
        )
        return max_row
    return ws.max_row


def validate_workbook(path: Path) -> dict[str, Any]:
    path = Path(path).resolve()
    wb = load_workbook(path, read_only=False, data_only=False)
    try:
        structural: list[str] = []
        missing_sheets = [name for name in REQUIRED_SHEETS if name not in wb.sheetnames]
        if missing_sheets:
            structural.append("Faltan hojas: " + ", ".join(missing_sheets))
            return {
                "workbook": str(path), "structural_errors": structural,
                "summary": {}, "issues": [],
            }
        products = wb["PRODUCTOS"]
        actual_headers = tuple(
            products.cell(HEADER_ROW, col).value
            for col in range(1, len(PRODUCT_HEADERS) + 1)
        )
        if actual_headers != PRODUCT_HEADERS:
            structural.append("Encabezados PRODUCTOS no coinciden con contrato canónico")
            return {
                "workbook": str(path), "structural_errors": structural,
                "summary": {}, "issues": [],
            }
        catalogs_ws = wb["CATALOGOS"]
        categories = _catalog_values(catalogs_ws, 1)
        units = _catalog_values(catalogs_ws, 5)
        providers = _catalog_values(catalogs_ws, 8)
        last_row = _product_last_row(products)
        records: list[tuple[int, dict[str, Any]]] = []
        issues_by_row: dict[int, list[ValidationIssue]] = {}
        staging_ids: dict[str, list[int]] = {}
        local_ids: dict[str, list[int]] = {}
        duplicate_keys: dict[tuple[str, str, str, str], list[int]] = {}
        barcodes: dict[str, list[int]] = {}

        for row in range(DATA_START_ROW, last_row + 1):
            record = {
                header: products.cell(row, col).value
                for col, header in enumerate(PRODUCT_HEADERS, 1)
            }
            records.append((row, record))
            row_issues: list[ValidationIssue] = []
            staging = str(record.get("registro_inventario_id") or "").strip()
            if not staging:
                row_issues.append(ValidationIssue(row, "STAGING_ID_MISSING", "registro_inventario_id", "Falta ID staging"))
            else:
                try:
                    uuid.UUID(staging)
                except (ValueError, AttributeError):
                    row_issues.append(ValidationIssue(row, "STAGING_ID_INVALID", "registro_inventario_id", "ID staging no es UUID"))
                staging_ids.setdefault(staging, []).append(row)
            active = is_active_record(record)
            if active:
                row_issues.extend(validate_record(
                    record, row=row, categories=categories, units=units,
                    providers=providers,
                ))
                local_id = str(record.get("producto_local_id") or "").strip()
                if local_id:
                    local_ids.setdefault(local_id, []).append(row)
                name = normalized_text(record.get("nombre"))
                if name:
                    key = (
                        name, normalized_text(record.get("marca")),
                        normalized_text(record.get("presentacion")),
                        normalized_text(record.get("unidad_medida")),
                    )
                    duplicate_keys.setdefault(key, []).append(row)
                barcode = str(record.get("codigo_barras") or "").strip()
                if barcode:
                    barcodes.setdefault(barcode.casefold(), []).append(row)
            issues_by_row[row] = row_issues

        for staging, rows in staging_ids.items():
            if staging and len(rows) > 1:
                for row in rows:
                    issues_by_row[row].append(ValidationIssue(row, "STAGING_ID_DUPLICATE", "registro_inventario_id", "ID staging duplicado"))
        for local_id, rows in local_ids.items():
            if len(rows) > 1:
                for row in rows:
                    issues_by_row[row].append(ValidationIssue(row, "LOCAL_ID_DUPLICATE", "producto_local_id", "local_id duplicado"))
        for rows in duplicate_keys.values():
            if len(rows) > 1:
                for row in rows:
                    issues_by_row[row].append(ValidationIssue(row, "POTENTIAL_DUPLICATE", "nombre", "REVISAR_POSIBLE_DUPLICADO", "WARNING"))
        for rows in barcodes.values():
            if len(rows) > 1:
                for row in rows:
                    issues_by_row[row].append(ValidationIssue(row, "BARCODE_DUPLICATE", "codigo_barras", "Barcode duplicado", "WARNING"))

        statuses: dict[int, str] = {}
        all_issues: list[ValidationIssue] = []
        active_count = 0
        for row, record in records:
            row_issues = issues_by_row[row]
            status = status_from_issues(record, row_issues)
            statuses[row] = status
            if status:
                active_count += 1
            all_issues.extend(row_issues)
        summary = {
            "product_rows": active_count,
            "LISTO_COMPLETO": sum(1 for value in statuses.values() if value == "LISTO_COMPLETO"),
            "LISTO_SIN_BARCODE": sum(1 for value in statuses.values() if value == "LISTO_SIN_BARCODE"),
            "INCOMPLETO": sum(1 for value in statuses.values() if value == "INCOMPLETO"),
            "REVISAR": sum(1 for value in statuses.values() if value == "REVISAR"),
            "issues": len(all_issues),
        }
        return {
            "workbook": str(path), "structural_errors": structural,
            "summary": summary,
            "issues": [asdict(issue) for issue in all_issues],
            "row_status": {str(row): status for row, status in statuses.items() if status},
        }
    finally:
        wb.close()


def main() -> int:
    parser = argparse.ArgumentParser(description="Validar workbook FERREPRO sin importar datos")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--json", type=Path, help="guardar reporte JSON")
    parser.add_argument("--csv", type=Path, help="guardar detalle CSV")
    args = parser.parse_args()
    result = validate_workbook(args.workbook)
    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        args.json.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    if args.csv:
        args.csv.parent.mkdir(parents=True, exist_ok=True)
        with args.csv.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=("row", "code", "field", "message", "kind"))
            writer.writeheader()
            writer.writerows(result["issues"])
    print("FERREPRO · validación de inventario")
    if result["structural_errors"]:
        for error in result["structural_errors"]:
            print("ERROR ESTRUCTURAL:", error)
        return 2
    for key, value in result["summary"].items():
        print(f"{key}: {value}")
    if result["issues"]:
        print("Primeros hallazgos:")
        for issue in result["issues"][:20]:
            print(f"  fila {issue['row']} [{issue['code']}] {issue['message']}")
    return 1 if result["summary"].get("REVISAR") else 0


if __name__ == "__main__":
    raise SystemExit(main())
