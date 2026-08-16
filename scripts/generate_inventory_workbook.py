# -*- coding: utf-8 -*-
"""Genera el workbook operativo de inventario físico FERREPRO.

Modos:
  --blank
  --from-db ferreteria.db  (SQLite se abre estrictamente read-only)
"""
from __future__ import annotations

import argparse
import sqlite3
import sys
from pathlib import Path
from typing import Any, Iterable, Mapping, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from product_inventory_contract import (
    BARCODE_STATUS_VALUES,
    CATEGORY_UI_VALUES,
    DATA_START_ROW,
    DEFAULT_BLANK_ROWS,
    FIELDS,
    HEADER_ROW,
    PRODUCT_FIELDS,
    PRODUCT_HEADERS,
    UNIT_UI_VALUES,
    YES_NO_VALUES,
    mapping_rows,
    stable_inventory_id,
)


DEFAULT_OUTPUT = REPO_ROOT / "plantillas" / "FERREPRO_Inventario_Maestro.xlsx"

NAVY = "17324D"
TEAL = "177E89"
GOLD = "E3B341"
LIGHT_BLUE = "E8F1F5"
LIGHT_YELLOW = "FFF4CC"
LIGHT_PURPLE = "F1E8FF"
LIGHT_GRAY = "EEF1F4"
LIGHT_GREEN = "E8F5E9"
WHITE = "FFFFFF"
RED = "FCE8E6"
TEXT = "243447"
THIN_GRAY = Side(style="thin", color="D4DCE3")


def _readonly_connection(path: Path) -> sqlite3.Connection:
    target = path.resolve()
    conn = sqlite3.connect(target.as_uri() + "?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA query_only = ON")
    return conn


def _columns(conn, table: str) -> set[str]:
    return {str(row[1]) for row in conn.execute(f"PRAGMA table_info({table})")}


def _unique_text(values: Iterable[Any]) -> list[str]:
    result: dict[str, str] = {}
    for value in values:
        text = str(value or "").strip()
        if text:
            result.setdefault(text.casefold(), text)
    return sorted(result.values(), key=str.casefold)


def load_source_data(db_path: Path) -> tuple[list[dict[str, Any]], dict[str, list[Any]]]:
    conn = _readonly_connection(db_path)
    try:
        product_cols = _columns(conn, "productos")
        provider_cols = _columns(conn, "proveedores")
        if not product_cols:
            raise ValueError("La base no contiene tabla productos")
        providers = []
        if provider_cols:
            providers = [
                {"id": int(row[0]), "nombre": str(row[1] or "").strip()}
                for row in conn.execute(
                    "SELECT id, nombre FROM proveedores "
                    "WHERE COALESCE(activo, 1)=1 ORDER BY LOWER(nombre), id"
                ).fetchall()
                if str(row[1] or "").strip()
            ]
        provider_by_id = {item["id"]: item["nombre"] for item in providers}
        wanted = [
            "id", "local_id", "codigo_barras", "nombre", "categoria", "marca",
            "presentacion", "proveedor_id", "precio_compra", "precio_venta",
            "stock_minimo", "unidad_medida", "permite_decimales", "viene_en_caja",
            "unidades_por_caja", "vende_por_empaque", "ubicacion", "descripcion",
            "stock", "activo",
        ]
        selected = [name for name in wanted if name in product_cols]
        rows = conn.execute(
            f"SELECT {', '.join(selected)} FROM productos ORDER BY id"
        ).fetchall()
        records: list[dict[str, Any]] = []
        categories, brands, units = set(CATEGORY_UI_VALUES), set(), set(UNIT_UI_VALUES)
        for row in rows:
            source = dict(row)
            local_id = str(source.get("local_id") or "").strip()
            record = {
                "registro_inventario_id": stable_inventory_id(
                    f"existing:{local_id or source.get('id')}"
                ),
                "tipo_registro": "EXISTENTE",
                "activo_sistema": "ACTIVO" if int(source.get("activo") or 0) else "INACTIVO",
                "producto_id_sistema": source.get("id"),
                "producto_local_id": local_id,
                "codigo_sistema_actual": source.get("codigo_barras"),
                "nombre": source.get("nombre"),
                "categoria": source.get("categoria"),
                "marca": source.get("marca"),
                "presentacion": source.get("presentacion"),
                "proveedor": provider_by_id.get(source.get("proveedor_id"), ""),
                "proveedor_id": source.get("proveedor_id"),
                "precio_compra": source.get("precio_compra"),
                "precio_venta": source.get("precio_venta"),
                "stock_minimo": source.get("stock_minimo", 10),
                "unidad_medida": source.get("unidad_medida") or "UNIDAD",
                "permite_decimales": "SI" if int(source.get("permite_decimales") or 0) else "NO",
                "viene_en_caja": "SI" if int(source.get("viene_en_caja") or 0) else "NO",
                "unidades_por_caja": source.get("unidades_por_caja") or 1,
                "vende_por_empaque": "SI" if int(source.get("vende_por_empaque") or 0) else "NO",
                "ubicacion": source.get("ubicacion"),
                "descripcion": source.get("descripcion"),
                "stock_sistema_actual": source.get("stock"),
                "cantidad_contada": None,
                "codigo_barras": None,
                "estado_barcode": "PENDIENTE",
            }
            records.append(record)
            if str(source.get("categoria") or "").strip():
                categories.add(str(source["categoria"]).strip())
            if str(source.get("marca") or "").strip():
                brands.add(str(source["marca"]).strip())
            if str(source.get("unidad_medida") or "").strip():
                units.add(str(source["unidad_medida"]).strip().upper())
        return records, {
            "categorias": _unique_text(categories),
            "marcas": _unique_text(brands),
            "unidades": _unique_text(units),
            "proveedores": providers,
        }
    finally:
        conn.close()


def blank_catalogs() -> dict[str, list[Any]]:
    return {
        "categorias": list(CATEGORY_UI_VALUES),
        "marcas": [],
        "unidades": list(UNIT_UI_VALUES),
        "proveedores": [],
    }


def _blank_record(slot: int, mode: str) -> dict[str, Any]:
    return {
        "registro_inventario_id": stable_inventory_id(f"{mode}:new-slot:{slot}"),
        "tipo_registro": "NUEVO",
        "activo_sistema": "ACTIVO",
        "stock_minimo": 10,
        "unidad_medida": "UNIDAD",
        "permite_decimales": "NO",
        "viene_en_caja": "NO",
        "unidades_por_caja": 1,
        "vende_por_empaque": "NO",
        "estado_barcode": "PENDIENTE",
    }


def _field_col(key: str) -> str:
    index = next(i for i, field in enumerate(PRODUCT_FIELDS, 1) if field.key == key)
    return get_column_letter(index)


def _active_formula(row: int) -> str:
    keys = (
        "nombre", "categoria", "marca", "presentacion", "proveedor",
        "precio_compra", "precio_venta", "ubicacion", "descripcion",
        "cantidad_contada", "pasillo", "estante", "codigo_barras",
        "observaciones_inventario",
    )
    refs = ",".join(f"{_field_col(key)}{row}" for key in keys)
    type_ref = f"{_field_col('tipo_registro')}{row}"
    return f'OR({type_ref}="EXISTENTE",COUNTA({refs})>0)'


def _row_formulas(row: int, last_row: int) -> dict[str, str]:
    c = {key: _field_col(key) for key in (
        "registro_inventario_id", "nombre", "categoria", "marca", "presentacion",
        "proveedor", "precio_compra", "precio_venta", "stock_minimo",
        "unidad_medida", "permite_decimales", "viene_en_caja",
        "unidades_por_caja", "stock_sistema_actual", "cantidad_contada",
        "codigo_barras", "estado_barcode", "errores_validacion",
        "posible_duplicado",
    )}
    active = _active_formula(row)
    required_missing = "OR(" + ",".join(
        f'{_field_col(key)}{row}=""'
        for key in (
            "nombre", "categoria", "precio_compra", "precio_venta",
            "stock_minimo", "unidad_medida", "cantidad_contada",
        )
    ) + ")"
    errors = (
        f'=IF(NOT({active}),"",TEXTJOIN(" | ",TRUE,'
        f'IF({c["registro_inventario_id"]}{row}="","Falta ID staging",""),'
        f'IF(COUNTIF(${c["registro_inventario_id"]}${DATA_START_ROW}:${c["registro_inventario_id"]}${last_row},{c["registro_inventario_id"]}{row})>1,"ID staging duplicado",""),'
        f'IF({c["nombre"]}{row}="","Falta nombre",""),'
        f'IF(LEN({c["nombre"]}{row})>150,"Nombre >150 caracteres",""),'
        f'IF({c["categoria"]}{row}="","Falta categoría",IF(COUNTIF(CategoriasLista,{c["categoria"]}{row})=0,"Categoría inválida","")),'
        f'IF({c["precio_compra"]}{row}="","Falta precio compra",IF(OR(NOT(ISNUMBER({c["precio_compra"]}{row})),{c["precio_compra"]}{row}<0),"Precio compra inválido","")),'
        f'IF({c["precio_venta"]}{row}="","Falta precio venta",IF(OR(NOT(ISNUMBER({c["precio_venta"]}{row})),{c["precio_venta"]}{row}<=0,{c["precio_venta"]}{row}<{c["precio_compra"]}{row}),"Precio venta inválido","")),'
        f'IF({c["stock_minimo"]}{row}="","Falta stock mínimo",IF(OR(NOT(ISNUMBER({c["stock_minimo"]}{row})),{c["stock_minimo"]}{row}<0,MOD({c["stock_minimo"]}{row},1)<>0),"Stock mínimo inválido","")),'
        f'IF({c["unidad_medida"]}{row}="","Falta unidad",IF(COUNTIF(UnidadesLista,{c["unidad_medida"]}{row})=0,"Unidad inválida","")),'
        f'IF(AND({c["proveedor"]}{row}<>"",COUNTIF(ProveedoresLista,{c["proveedor"]}{row})=0),"Proveedor no catalogado",""),'
        f'IF({c["cantidad_contada"]}{row}="","Falta cantidad contada",IF(OR(NOT(ISNUMBER({c["cantidad_contada"]}{row})),{c["cantidad_contada"]}{row}<0,ABS({c["cantidad_contada"]}{row}*1000-ROUND({c["cantidad_contada"]}{row}*1000,0))>0.0000001),"Cantidad inválida",IF(AND({c["permite_decimales"]}{row}="NO",MOD({c["cantidad_contada"]}{row},1)<>0),"Cantidad fraccionaria no permitida",""))),'
        f'IF(AND({c["viene_en_caja"]}{row}="SI",OR(NOT(ISNUMBER({c["unidades_por_caja"]}{row})),{c["unidades_por_caja"]}{row}<1,MOD({c["unidades_por_caja"]}{row},1)<>0)),"Empaque inválido",""),'
        f'IF(OR(AND({c["codigo_barras"]}{row}="",{c["estado_barcode"]}{row}="ESCANEADO"),AND({c["codigo_barras"]}{row}<>"",{c["estado_barcode"]}{row}<>"ESCANEADO")),"Barcode/estado inconsistente",""),'
        f'IF({c["posible_duplicado"]}{row}<>"",{c["posible_duplicado"]}{row},"")))'
    )
    duplicate = (
        f'=IF({c["nombre"]}{row}="","",IF(OR('
        f'COUNTIFS(${c["nombre"]}${DATA_START_ROW}:${c["nombre"]}${last_row},{c["nombre"]}{row},'
        f'${c["marca"]}${DATA_START_ROW}:${c["marca"]}${last_row},{c["marca"]}{row},'
        f'${c["presentacion"]}${DATA_START_ROW}:${c["presentacion"]}${last_row},{c["presentacion"]}{row},'
        f'${c["unidad_medida"]}${DATA_START_ROW}:${c["unidad_medida"]}${last_row},{c["unidad_medida"]}{row})>1,'
        f'AND({c["codigo_barras"]}{row}<>"",COUNTIF(${c["codigo_barras"]}${DATA_START_ROW}:${c["codigo_barras"]}${last_row},{c["codigo_barras"]}{row})>1)),'
        '"REVISAR_POSIBLE_DUPLICADO",""))'
    )
    status = (
        f'=IF(NOT({active}),"",IF({c["posible_duplicado"]}{row}<>"","REVISAR",'
        f'IF({required_missing},"INCOMPLETO",IF({c["errores_validacion"]}{row}<>"","REVISAR",'
        f'IF(AND({c["codigo_barras"]}{row}<>"",{c["estado_barcode"]}{row}="ESCANEADO"),"LISTO_COMPLETO","LISTO_SIN_BARCODE")))))'
    )
    return {
        "diferencia_unidades": f'=IF({_field_col("cantidad_contada")}{row}="","",{_field_col("cantidad_contada")}{row}-IF({_field_col("stock_sistema_actual")}{row}="",0,{_field_col("stock_sistema_actual")}{row}))',
        "valor_inventario_costo": f'=IF(OR({_field_col("cantidad_contada")}{row}="",{_field_col("precio_compra")}{row}=""),"",{_field_col("cantidad_contada")}{row}*{_field_col("precio_compra")}{row})',
        "valor_potencial_venta": f'=IF(OR({_field_col("cantidad_contada")}{row}="",{_field_col("precio_venta")}{row}=""),"",{_field_col("cantidad_contada")}{row}*{_field_col("precio_venta")}{row})',
        "margen_unitario": f'=IF(OR({_field_col("precio_compra")}{row}="",{_field_col("precio_venta")}{row}=""),"",{_field_col("precio_venta")}{row}-{_field_col("precio_compra")}{row})',
        "margen_porcentaje": f'=IFERROR(IF({_field_col("precio_compra")}{row}>0,({_field_col("precio_venta")}{row}-{_field_col("precio_compra")}{row})/{_field_col("precio_compra")}{row},""),"")',
        "posible_duplicado": duplicate,
        "errores_validacion": errors,
        "estado_registro": status,
    }


def _title(ws, title: str, subtitle: str, last_col: int) -> None:
    end = get_column_letter(last_col)
    ws.merge_cells(f"A1:{end}1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells(f"A2:{end}2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Aptos", size=10, color=TEXT)
    ws["A2"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 28


def _instructions_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("INSTRUCCIONES")
    ws.sheet_view.showGridLines = False
    _title(ws, "FERREPRO · Inventario físico maestro", "Complete primero el catálogo y conteo. El barcode puede escanearse en una fase posterior.", 8)
    sections = [
        (4, "Cómo trabajar", [
            "Una fila representa un producto/SKU. No fusione variantes distintas.",
            "No cambie encabezados ni borre registro_inventario_id o IDs técnicos.",
            "Escriba el conteo físico en cantidad_contada; stock_sistema_actual es solo referencia.",
            "Los precios y cantidades deben ser celdas numéricas, no texto.",
            "cantidad_contada admite hasta 3 decimales; active permite_decimales para fracciones.",
            "codigo_barras puede quedar vacío y estado_barcode debe permanecer PENDIENTE.",
            "No duplique una fila porque un SKU tenga varios barcodes; códigos adicionales se modelarán después.",
            "Use observaciones_inventario para presentaciones, ubicación o dudas.",
        ]),
        (15, "Flujo posterior", [
            "Excel completo → LISTO_SIN_BARCODE → fase de escaneo → asociar scan al registro → LISTO_COMPLETO → importación controlada.",
            "Este archivo no importa productos, no modifica stock y no genera códigos FRP/EAN.",
        ]),
    ]
    for start, heading, bullets in sections:
        ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=8)
        cell = ws.cell(start, 1, heading)
        cell.font = Font(name="Aptos", size=13, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        for offset, bullet in enumerate(bullets, 1):
            ws.merge_cells(start_row=start + offset, start_column=1, end_row=start + offset, end_column=8)
            item = ws.cell(start + offset, 1, "• " + bullet)
            item.font = Font(name="Aptos", size=11, color=TEXT)
            item.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[start + offset].height = 28
    legend_row = 20
    ws.merge_cells(start_row=legend_row, start_column=1, end_row=legend_row, end_column=8)
    ws.cell(legend_row, 1, "Leyenda").font = Font(bold=True, size=13, color=WHITE)
    ws.cell(legend_row, 1).fill = PatternFill("solid", fgColor=NAVY)
    legends = [
        ("OBLIGATORIO", LIGHT_YELLOW, "Debe existir para LISTO_SIN_BARCODE"),
        ("OPCIONAL", WHITE, "Mejora la ficha pero no bloquea"),
        ("AUTOMÁTICO", LIGHT_GRAY, "Fórmula o identidad técnica; no editar"),
        ("PENDIENTE BARCODE", LIGHT_PURPLE, "Se completa al escanear físicamente"),
    ]
    for i, (label, color, desc) in enumerate(legends, legend_row + 1):
        ws.cell(i, 1, label).font = Font(bold=True, color=TEXT)
        ws.cell(i, 1).fill = PatternFill("solid", fgColor=color)
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=8)
        ws.cell(i, 2, desc)
    for col in range(1, 9):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.freeze_panes = "A3"


def _catalogs_sheet(wb: Workbook, catalogs: Mapping[str, list[Any]]) -> None:
    ws = wb.create_sheet("CATALOGOS")
    ws.sheet_view.showGridLines = False
    _title(ws, "Catálogos para captura", "Categoría y unidad son listas controladas. Marca es abierta. Proveedor es opcional y debe resolverse si es nuevo.", 8)
    headers = {1: "CATEGORIAS", 3: "MARCAS_EXISTENTES", 5: "UNIDADES", 7: "PROVEEDOR_ID", 8: "PROVEEDORES"}
    for col, value in headers.items():
        cell = ws.cell(4, col, value)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TEAL)
    for row, value in enumerate(catalogs["categorias"], 5):
        ws.cell(row, 1, value)
    for row, value in enumerate(catalogs["marcas"], 5):
        ws.cell(row, 3, value)
    for row, value in enumerate(catalogs["unidades"], 5):
        ws.cell(row, 5, value)
    for row, item in enumerate(catalogs["proveedores"], 5):
        ws.cell(row, 7, item["id"])
        ws.cell(row, 8, item["nombre"])
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["G"].width = 16
    ws.column_dimensions["H"].width = 34
    ws.freeze_panes = "A5"
    cat_end = max(5, 4 + len(catalogs["categorias"]))
    unit_end = max(5, 4 + len(catalogs["unidades"]))
    prov_end = max(5, 4 + len(catalogs["proveedores"]))
    wb.defined_names.add(DefinedName("CategoriasLista", attr_text=f"'CATALOGOS'!$A$5:$A${cat_end}"))
    wb.defined_names.add(DefinedName("UnidadesLista", attr_text=f"'CATALOGOS'!$E$5:$E${unit_end}"))
    wb.defined_names.add(DefinedName("ProveedoresLista", attr_text=f"'CATALOGOS'!$H$5:$H${prov_end}"))


def _products_sheet(
    wb: Workbook, records: list[dict[str, Any]], blank_rows: int, mode: str
) -> tuple[int, int]:
    ws = wb.create_sheet("PRODUCTOS")
    ws.sheet_view.showGridLines = False
    last_col = len(PRODUCT_FIELDS)
    _title(ws, "PRODUCTOS · captura maestra", "Amarillo = obligatorio · gris = automático · morado = barcode posterior. cantidad_contada nunca se prellena desde stock.", last_col)
    block_colors = {
        "IDENTIFICACIÓN": NAVY,
        "INFORMACIÓN COMERCIAL": TEAL,
        "INVENTARIO FÍSICO": "4D7C8A",
        "CÁLCULOS": "637083",
        "BARCODE": "7353A6",
        "CONTROL": "8A6D3B",
    }
    start = 1
    current = PRODUCT_FIELDS[0].block
    for index, field in enumerate(PRODUCT_FIELDS, 1):
        if field.block != current:
            ws.merge_cells(start_row=3, start_column=start, end_row=3, end_column=index - 1)
            cell = ws.cell(3, start, current)
            cell.fill = PatternFill("solid", fgColor=block_colors[current])
            cell.font = Font(bold=True, color=WHITE)
            cell.alignment = Alignment(horizontal="center")
            start, current = index, field.block
    ws.merge_cells(start_row=3, start_column=start, end_row=3, end_column=last_col)
    ws.cell(3, start, current).fill = PatternFill("solid", fgColor=block_colors[current])
    ws.cell(3, start).font = Font(bold=True, color=WHITE)
    ws.cell(3, start).alignment = Alignment(horizontal="center")
    for col, field in enumerate(PRODUCT_FIELDS, 1):
        cell = ws.cell(HEADER_ROW, col, field.excel_column)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY)
        ws.column_dimensions[get_column_letter(col)].width = field.width
    ws.row_dimensions[3].height = 23
    ws.row_dimensions[4].height = 42

    all_records = list(records) + [
        _blank_record(slot, mode) for slot in range(1, blank_rows + 1)
    ]
    last_row = DATA_START_ROW + len(all_records) - 1
    formula_fields = {field.key for field in PRODUCT_FIELDS if field.data_type == "FORMULA"}
    for offset, record in enumerate(all_records):
        row = DATA_START_ROW + offset
        formulas = _row_formulas(row, last_row)
        for col, field in enumerate(PRODUCT_FIELDS, 1):
            cell = ws.cell(row, col)
            if field.key in formula_fields:
                cell.value = formulas[field.key]
            else:
                cell.value = record.get(field.key, field.default)
            cell.font = Font(name="Aptos", size=10, color=TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=field.width >= 25)
            cell.number_format = field.number_format
            cell.border = Border(bottom=Side(style="hair", color="E3E8EC"))
            if field.generated_by_system:
                cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)
                cell.protection = Protection(locked=True)
            elif field.barcode_related:
                cell.fill = PatternFill("solid", fgColor=LIGHT_PURPLE)
                cell.protection = Protection(locked=False)
            elif field.required:
                cell.fill = PatternFill("solid", fgColor=LIGHT_YELLOW)
                cell.protection = Protection(locked=False)
            else:
                cell.protection = Protection(locked=False)

    table = Table(displayName="ProductosInventario", ref=f"A{HEADER_ROW}:{get_column_letter(last_col)}{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)
    ws.freeze_panes = "F5"
    ws.auto_filter.ref = table.ref

    validations = []
    for key, formula in (
        ("categoria", "=CategoriasLista"),
        ("unidad_medida", "=UnidadesLista"),
        ("proveedor", "=ProveedoresLista"),
    ):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.errorStyle = "warning" if key == "proveedor" else "stop"
        dv.error = "Seleccione un valor de CATALOGOS."
        dv.errorTitle = "Catálogo FERREPRO"
        dv.showErrorMessage = True
        validations.append((dv, key))
    for key in ("permite_decimales", "viene_en_caja", "vende_por_empaque"):
        dv = DataValidation(type="list", formula1='"NO,SI"', allow_blank=False)
        validations.append((dv, key))
    barcode_dv = DataValidation(type="list", formula1='"PENDIENTE,ESCANEADO"', allow_blank=False)
    validations.append((barcode_dv, "estado_barcode"))
    for dv, key in validations:
        ws.add_data_validation(dv)
        col = _field_col(key)
        dv.add(f"{col}{DATA_START_ROW}:{col}{last_row}")

    count_col = _field_col("cantidad_contada")
    decimal_dv = DataValidation(
        type="custom",
        formula1=f'=OR({count_col}{DATA_START_ROW}="",AND(ISNUMBER({count_col}{DATA_START_ROW}),{count_col}{DATA_START_ROW}>=0,ABS({count_col}{DATA_START_ROW}*1000-ROUND({count_col}{DATA_START_ROW}*1000,0))<0.0000001))',
        allow_blank=True,
    )
    decimal_dv.error = "Use un número no negativo con máximo 3 decimales."
    decimal_dv.errorTitle = "Cantidad inválida"
    decimal_dv.showErrorMessage = True
    ws.add_data_validation(decimal_dv)
    decimal_dv.add(f"{count_col}{DATA_START_ROW}:{count_col}{last_row}")

    barcode_col = _field_col("codigo_barras")
    barcode_unique = DataValidation(
        type="custom",
        formula1=f'=OR({barcode_col}{DATA_START_ROW}="",COUNTIF(${barcode_col}${DATA_START_ROW}:${barcode_col}${last_row},{barcode_col}{DATA_START_ROW})=1)',
        allow_blank=True,
    )
    barcode_unique.errorStyle = "warning"
    barcode_unique.error = "Barcode repetido: revise, no fusione automáticamente."
    barcode_unique.showErrorMessage = True
    ws.add_data_validation(barcode_unique)
    barcode_unique.add(f"{barcode_col}{DATA_START_ROW}:{barcode_col}{last_row}")

    state_col = _field_col("estado_registro")
    state_range = f"{state_col}{DATA_START_ROW}:{state_col}{last_row}"
    ws.conditional_formatting.add(state_range, FormulaRule(formula=[f'{state_col}{DATA_START_ROW}="LISTO_COMPLETO"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(state_range, FormulaRule(formula=[f'{state_col}{DATA_START_ROW}="LISTO_SIN_BARCODE"'], fill=PatternFill("solid", fgColor="DDEBF7")))
    ws.conditional_formatting.add(state_range, FormulaRule(formula=[f'{state_col}{DATA_START_ROW}="REVISAR"'], fill=PatternFill("solid", fgColor=RED)))
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return last_row, len(records)


def _summary_sheet(wb: Workbook, last_row: int) -> None:
    ws = wb.create_sheet("RESUMEN")
    ws.sheet_view.showGridLines = False
    _title(ws, "Resumen del inventario", "Indicadores calculados desde PRODUCTOS. Se actualizan al abrir/recalcular el archivo en Excel.", 6)
    state = _field_col("estado_registro")
    quantity = _field_col("cantidad_contada")
    cost_value = _field_col("valor_inventario_costo")
    sale_value = _field_col("valor_potencial_venta")
    barcode_state = _field_col("estado_barcode")
    record_type = _field_col("tipo_registro")
    name = _field_col("nombre")
    metrics = [
        ("Total filas de productos", f'=COUNTIF(\'PRODUCTOS\'!${name}${DATA_START_ROW}:${name}${last_row},"<>")', "#,##0"),
        ("Productos completos", f'=COUNTIF(\'PRODUCTOS\'!${state}${DATA_START_ROW}:${state}${last_row},"LISTO_COMPLETO")', "#,##0"),
        ("LISTO_SIN_BARCODE", f'=COUNTIF(\'PRODUCTOS\'!${state}${DATA_START_ROW}:${state}${last_row},"LISTO_SIN_BARCODE")', "#,##0"),
        ("Productos incompletos", f'=COUNTIF(\'PRODUCTOS\'!${state}${DATA_START_ROW}:${state}${last_row},"INCOMPLETO")', "#,##0"),
        ("Barcodes pendientes", f'=COUNTIFS(\'PRODUCTOS\'!${barcode_state}${DATA_START_ROW}:${barcode_state}${last_row},"PENDIENTE",\'PRODUCTOS\'!${name}${DATA_START_ROW}:${name}${last_row},"<>")', "#,##0"),
        ("Productos a revisar", f'=COUNTIF(\'PRODUCTOS\'!${state}${DATA_START_ROW}:${state}${last_row},"REVISAR")', "#,##0"),
        ("Cantidad total contada", f'=SUM(\'PRODUCTOS\'!${quantity}${DATA_START_ROW}:${quantity}${last_row})', "#,##0.000"),
        ("Valor total a costo", f'=SUM(\'PRODUCTOS\'!${cost_value}${DATA_START_ROW}:${cost_value}${last_row})', '"$"#,##0.00'),
        ("Valor potencial de venta", f'=SUM(\'PRODUCTOS\'!${sale_value}${DATA_START_ROW}:${sale_value}${last_row})', '"$"#,##0.00'),
        ("Productos nuevos", f'=COUNTIFS(\'PRODUCTOS\'!${record_type}${DATA_START_ROW}:${record_type}${last_row},"NUEVO",\'PRODUCTOS\'!${name}${DATA_START_ROW}:${name}${last_row},"<>")', "#,##0"),
        ("Productos existentes", f'=COUNTIFS(\'PRODUCTOS\'!${record_type}${DATA_START_ROW}:${record_type}${last_row},"EXISTENTE",\'PRODUCTOS\'!${name}${DATA_START_ROW}:${name}${last_row},"<>")', "#,##0"),
    ]
    for index, (label, formula, fmt) in enumerate(metrics):
        row = 4 + index
        ws.cell(row, 1, label)
        ws.cell(row, 1).font = Font(bold=True, color=TEXT)
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(row, 2, formula)
        ws.cell(row, 2).font = Font(size=14, bold=True, color=NAVY)
        ws.cell(row, 2).number_format = fmt
        ws.cell(row, 2).fill = PatternFill("solid", fgColor=WHITE)
        for col in (1, 2):
            ws.cell(row, col).border = Border(bottom=THIN_GRAY)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 24
    ws.freeze_panes = "A4"


def _mapping_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("MAPEO_FERREPRO")
    ws.sheet_view.showGridLines = False
    _title(ws, "Contrato de mapeo FERREPRO", "Fuente técnica para el futuro importador. No cambiar sin actualizar product_inventory_contract.py y tests.", 9)
    headers = ("Excel", "Campo FERREPRO", "Tabla", "Columna DB", "Required", "Tipo", "Default", "Transformación", "Observación")
    for col, value in enumerate(headers, 1):
        cell = ws.cell(4, col, value)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True)
    for row_index, values in enumerate(mapping_rows(), 5):
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_index, col, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color="E3E8EC"))
    widths = (28, 28, 18, 34, 12, 18, 18, 35, 55)
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    end = 4 + len(mapping_rows())
    table = Table(displayName="MapeoFerrepro", ref=f"A4:I{end}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A5"


def _barcode_sheet(wb: Workbook, last_row: int) -> None:
    ws = wb.create_sheet("BARCODES_PENDIENTES")
    ws.sheet_view.showGridLines = False
    headers = (
        "registro_inventario_id", "nombre", "marca", "presentacion",
        "cantidad_contada", "codigo_barras", "estado_barcode",
    )
    _title(ws, "Barcodes pendientes", "Vista para la futura fase DIG-X6266. No implementa scanner ni persistencia.", len(headers))
    for col, value in enumerate(headers, 1):
        cell = ws.cell(4, col, value)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor="7353A6")
        cell.alignment = Alignment(wrap_text=True)
    for source_row in range(DATA_START_ROW, last_row + 1):
        target_row = source_row
        active_condition = f'\'PRODUCTOS\'!${_field_col("nombre")}{source_row}<>""'
        for col, key in enumerate(headers, 1):
            source_col = _field_col(key)
            ws.cell(
                target_row,
                col,
                f'=IFERROR(IF({active_condition},\'PRODUCTOS\'!${source_col}{source_row},""),"")',
            )
            ws.cell(target_row, col).fill = PatternFill("solid", fgColor=LIGHT_PURPLE if key in ("codigo_barras", "estado_barcode") else WHITE)
    widths = (38, 32, 22, 24, 19, 25, 18)
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.auto_filter.ref = f"A4:G{last_row}"
    ws.freeze_panes = "A5"


def generate_workbook(
    output: Path, *, db_path: Optional[Path] = None,
    blank_rows: int = DEFAULT_BLANK_ROWS,
) -> dict[str, Any]:
    if blank_rows < 1:
        raise ValueError("blank_rows debe ser >= 1")
    if db_path is None:
        mode, records, catalogs = "blank", [], blank_catalogs()
    else:
        mode = "from-db"
        records, catalogs = load_source_data(Path(db_path))
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = "FERREPRO Inventario Maestro"
    wb.properties.subject = "Inventario físico y staging de barcodes"
    wb.properties.creator = "FERREPRO"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    _instructions_sheet(wb)
    last_row, existing_count = _products_sheet(wb, records, blank_rows, mode)
    _catalogs_sheet(wb, catalogs)
    _summary_sheet(wb, last_row)
    _mapping_sheet(wb)
    _barcode_sheet(wb, last_row)
    wb._sheets = [
        wb["INSTRUCCIONES"], wb["PRODUCTOS"], wb["CATALOGOS"],
        wb["RESUMEN"], wb["MAPEO_FERREPRO"], wb["BARCODES_PENDIENTES"],
    ]
    output = Path(output).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    # Prueba inmediata de que el ZIP/OpenXML puede reabrirse.
    check = load_workbook(output, read_only=False, data_only=False)
    try:
        if check.sheetnames != [
            "INSTRUCCIONES", "PRODUCTOS", "CATALOGOS", "RESUMEN",
            "MAPEO_FERREPRO", "BARCODES_PENDIENTES",
        ]:
            raise RuntimeError("Workbook generado con hojas inesperadas")
    finally:
        check.close()
    return {
        "output": str(output), "mode": mode, "existing_products": existing_count,
        "blank_slots": blank_rows, "product_rows": existing_count + blank_rows,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Generar FERREPRO Inventario Maestro")
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--blank", action="store_true", help="plantilla sin productos")
    mode.add_argument("--from-db", type=Path, help="SQLite legacy en modo read-only")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--blank-rows", type=int, default=DEFAULT_BLANK_ROWS)
    args = parser.parse_args()
    result = generate_workbook(
        args.output,
        db_path=args.from_db if args.from_db else None,
        blank_rows=args.blank_rows,
    )
    print(
        f"OK {result['mode']}: {result['existing_products']} existentes + "
        f"{result['blank_slots']} espacios nuevos -> {result['output']}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
